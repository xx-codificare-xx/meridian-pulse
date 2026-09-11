import streamlit as st
import sys
import os
import io
import json
import datetime

sys.path.append(os.path.dirname(os.path.abspath(__file__)))

from config import (
    APP_TITLE, APP_TAGLINE,
    TAGS, MAX_DASHBOARD_ARTICLES,
)
from transcripts.transcript_handler import read_transcript, get_transcript_files

# ─────────────────────────────────────────
# PAGE CONFIG
# ─────────────────────────────────────────
st.set_page_config(
    page_title=APP_TITLE,
    page_icon="M",
    layout="wide",
    #initial_sidebar_state="collapsed",
    initial_sidebar_state="expanded",

)

ARTICLES_PER_PAGE = 15

# ─────────────────────────────────────────
# STYLES
# ─────────────────────────────────────────
def load_styles(dark_mode: bool):
    bg          = "#0F1117" if dark_mode else "#F8F5F0"
    card_bg     = "#1C1E26" if dark_mode else "#FFFFFF"
    text_primary= "#F0EEE9" if dark_mode else "#1A1A1A"
    text_sec    = "#9B9B9B" if dark_mode else "#6B6B6B"
    accent      = "#4A90D9" if dark_mode else "#1B4F8A"
    border      = "#2E3140" if dark_mode else "#E8E4DC"
    tab_bg      = "#1C1E26" if dark_mode else "#FFFFFF"
    chip_bg     = "#2E3140" if dark_mode else "#EEF2F7"
    user_bubble = "#2E3140" if dark_mode else "#EEF2F7"
    bot_bubble  = "#1C1E26" if dark_mode else "#FFFFFF"

    st.markdown(f"""
    <style>
    @import url('https://fonts.googleapis.com/css2?family=Playfair+Display:wght@600;700&family=Inter:wght@300;400;500;600&display=swap');

    html, body, [class*="css"] {{
        background-color: {bg} !important;
        color: {text_primary} !important;
        font-family: 'Inter', sans-serif;
    }}

    #MainMenu, footer, header {{ visibility: hidden; }}
    .block-container {{ padding-top: 0rem !important; }}

    @keyframes fadeSlideUp {{
        from {{ opacity: 0; transform: translateY(30px); }}
        to   {{ opacity: 1; transform: translateY(0); }}
    }}
    @keyframes pulse {{
        0%, 100% {{ opacity: 1; }}
        50%       {{ opacity: 0.6; }}
    }}
    @keyframes gradientShift {{
        0%   {{ background-position: 0% 50%; }}
        50%  {{ background-position: 100% 50%; }}
        100% {{ background-position: 0% 50%; }}
    }}

    .chat-bubble-user{{
        background: #EEF2F7;
        border-radius: 1rem 1rem 0.2rem 1rem;
        padding: 0.7rem 1rem;
        margin-bottom: 0.5rem;
        font-size: 0.88rem;
        text-align: right;
    }}
    .chat-bubble-bot{{
        background: #FFFFFF;
        border: 1px solid #E8E4DC;
        border-left: 3px solid #1B4F8A;
        border-radius: 1rem 1rem 1rem 0.2rem;
        padding: 0.7rem 1rem;
        margin-bottom: 0.5rem;
        font-size: 0.88rem;
        line-height: 1.6;
    }}

    .hero {{
        background: linear-gradient(-45deg,
            {'#0a0e1a, #1a2340, #0d1b2a, #162032' if dark_mode else '#e8f4f8, #f0e8f8, #e8f0f8, #f8f0e8'});
        background-size: 400% 400%;
        animation: gradientShift 8s ease infinite;
        padding: 3rem 2rem 2rem;
        border-radius: 0 0 2rem 2rem;
        margin-bottom: 2rem;
        text-align: center;
    }}
    .hero-title {{
        font-family: 'Playfair Display', serif;
        font-size: 3.5rem;
        font-weight: 700;
        color: {accent};
        animation: fadeSlideUp 1s ease forwards;
        margin-bottom: 0.5rem;
        letter-spacing: -1px;
    }}
    .hero-tagline {{
        font-size: 1rem;
        color: {text_sec};
        font-weight: 300;
        animation: fadeSlideUp 1.2s ease forwards;
        letter-spacing: 3px;
        text-transform: uppercase;
    }}

    .stTabs [data-baseweb="tab-list"] {{
        background: {tab_bg};
        border-radius: 1rem;
        padding: 0.3rem;
        gap: 0.2rem;
        border: 1px solid {border};
        margin: 0 1rem 1.5rem;
    }}
    .stTabs [data-baseweb="tab"] {{
        border-radius: 0.7rem;
        font-family: 'Inter', sans-serif;
        font-weight: 500;
        font-size: 0.9rem;
        color: {text_sec};
        padding: 0.6rem 1.5rem;
        transition: all 0.3s ease;
    }}
    .stTabs [aria-selected="true"] {{
        background: {accent} !important;
        color: white !important;
    }}

    .news-card {{
        background: {card_bg};
        border: 1px solid {border};
        border-radius: 1rem;
        padding: 1.5rem;
        margin-bottom: 1rem;
        animation: fadeSlideUp 0.6s ease forwards;
        transition: transform 0.2s ease, box-shadow 0.2s ease;
    }}
    .news-card:hover {{
        transform: translateY(-3px);
        box-shadow: 0 8px 30px rgba(0,0,0,0.12);
    }}

    .badge {{
        display: inline-block;
        padding: 0.25rem 0.8rem;
        border-radius: 2rem;
        font-size: 0.72rem;
        font-weight: 600;
        margin-bottom: 0.8rem;
        letter-spacing: 0.8px;
        text-transform: uppercase;
    }}
    .badge-high    {{ background: #FFF0F0; color: #CC0000; border: 1px solid #FFB3B3; }}
    .badge-medium  {{ background: #FFFBF0; color: #B8860B; border: 1px solid #FFE082; }}
    .badge-low     {{ background: #F0FFF4; color: #276749; border: 1px solid #A8D5B5; }}
    .badge-minimal {{ background: #F5F5F5; color: #555555; border: 1px solid #CCCCCC; }}
    .pulse-high    {{ animation: pulse 2s ease-in-out infinite; }}

    .card-title {{
        font-family: 'Playfair Display', serif;
        font-size: 1.1rem;
        font-weight: 600;
        color: {text_primary};
        margin-bottom: 0.5rem;
        line-height: 1.4;
    }}
    .card-meta {{
        font-size: 0.8rem;
        color: {text_sec};
        margin-bottom: 0.8rem;
    }}
    .card-reasoning {{
        font-size: 0.88rem;
        color: {text_sec};
        line-height: 1.6;
        border-left: 3px solid {accent};
        padding-left: 0.8rem;
        margin-top: 0.8rem;
    }}
    .score-bar-bg {{
        background: {border};
        border-radius: 1rem;
        height: 6px;
        margin-top: 0.8rem;
    }}
    .score-bar-fill {{
        height: 6px;
        border-radius: 1rem;
        background: linear-gradient(90deg, {accent}, #E74C3C);
    }}
    .tag-chip {{
        display: inline-block;
        background: {chip_bg};
        color: {accent};
        border-radius: 1rem;
        padding: 0.2rem 0.7rem;
        font-size: 0.75rem;
        margin: 0.2rem;
        font-weight: 500;
    }}
    .divider {{
        height: 1px;
        background: {border};
        margin: 1.5rem 0;
    }}
    .section-label {{
        font-size: 0.75rem;
        font-weight: 600;
        letter-spacing: 1px;
        text-transform: uppercase;
        color: {accent};
        margin-bottom: 0.3rem;
    }}
    .hire-card {{
        background: {card_bg};
        border: 1px solid {border};
        border-radius: 1.5rem;
        padding: 3rem;
        text-align: center;
        max-width: 600px;
        margin: 2rem auto;
        animation: fadeSlideUp 0.8s ease forwards;
    }}
    .hire-name {{
        font-family: 'Playfair Display', serif;
        font-size: 2rem;
        font-weight: 700;
        color: {accent};
        margin-bottom: 0.5rem;
    }}
    .hire-role {{
        font-size: 0.9rem;
        color: {text_sec};
        margin-bottom: 2rem;
        letter-spacing: 2px;
        text-transform: uppercase;
    }}
    .contact-link {{
        display: block;
        color: {accent};
        text-decoration: none;
        font-size: 1rem;
        margin: 0.5rem 0;
        font-weight: 500;
    }}

    .chat-float {{
        position: fixed;
        bottom: 2rem;
        right: 2rem;
        width: 56px;
        height: 56px;
        background: {accent};
        border-radius: 50%;
        display: flex;
        align-items: center;
        justify-content: center;
        cursor: pointer;
        z-index: 9999;
        box-shadow: 0 4px 20px rgba(0,0,0,0.25);
        transition: transform 0.2s ease, box-shadow 0.2s ease;
    }}
    .chat-float:hover {{
        transform: scale(1.1);
        box-shadow: 0 6px 25px rgba(0,0,0,0.35);
    }}
    .chat-float svg {{
        width: 24px;
        height: 24px;
        fill: white;
    }}

    .chat-bubble-user {{
        background: {user_bubble};
        border-radius: 1rem 1rem 0.2rem 1rem;
        padding: 0.7rem 1rem;
        margin-bottom: 0.5rem;
        font-size: 0.88rem;
        text-align: right;
        color: {text_primary};
    }}
    .chat-bubble-bot {{
        background: {bot_bubble};
        border: 1px solid {border};
        border-left: 3px solid {accent};
        border-radius: 1rem 1rem 1rem 0.2rem;
        padding: 0.7rem 1rem;
        margin-bottom: 0.5rem;
        font-size: 0.88rem;
        line-height: 1.6;
        color: {text_primary};
    }}
    </style>
    """, unsafe_allow_html=True)
    
    st.markdown("""
    <div class="chat-float" onclick="
        const toggle = window.parent.document.querySelector('[data-testid=collapsedControl]');
        if(toggle) toggle.click();
        ">
        <svg viewBox="0 0 24 24" xmlns="http://www.w3.org/2000/svg">
            <path d="M20 2H4c-1.1 0-2 .9-2 2v18l4-4h14c1.1 0 2-.9 2-2V4c0-1.1-.9-2-2-2zm-2 12H6v-2h12v2zm0-3H6V9h12v2zm0-3H6V6h12v2z"/>
        </svg>
    </div>
    """, unsafe_allow_html=True)



# ─────────────────────────────────────────
# HERO
# ─────────────────────────────────────────
def render_hero():
    import base64
    logo_path = os.path.join(os.path.dirname(__file__), "assets", "logo.png")
    try:
        with open(logo_path, "rb") as f:
            logo_b64 = base64.b64encode(f.read()).decode()
        logo_src = f"data:image/png;base64,{logo_b64}"
    except FileNotFoundError:
        logo_src = "https://via.placeholder.com/64x64/1B4F8A/FFFFFF?text=M"

    st.markdown(f"""
    <div class="hero">
        <img src="{logo_src}"
             style="border-radius:100%;width:120px;height:120px;
                    margin-bottom:1rem;display:block;
                    margin-left:auto;margin-right:auto;"/>
        <div class="hero-title">{APP_TITLE}</div>
        <div class="hero-tagline">{APP_TAGLINE}</div>
    </div>
    """, unsafe_allow_html=True)



# ─────────────────────────────────────────
# BADGE
# ─────────────────────────────────────────
def get_badge(label: str) -> str:
    mapping = {
        "🔴 High":    ("high",    "High Priority"),
        "🟡 Medium":  ("medium",  "Medium"),
        "🟢 Low":     ("low",     "Low"),
        "🔵 Minimal": ("minimal", "Minimal"),
        "Unscored":   ("minimal", "Unscored"),
    }
    cls, text = mapping.get(label, ("minimal", label))
    pulse = " pulse-high" if cls == "high" else ""
    return f'<span class="badge badge-{cls}{pulse}">{text}</span>'


# ─────────────────────────────────────────
# NEWS CARD
# ─────────────────────────────────────────
def render_card(article: dict, idx: int):
    badge     = get_badge(article.get("urgency_label", "Unscored"))
    score     = article.get("urgency_score", 0.0)
    score_pct = int(score * 100)
    tags      = article.get("tags_matched", [])
    reasoning = article.get("reasoning", "")
    source    = article.get("source", "")
    published = article.get("published", "")
    title     = article.get("title", "Untitled")
    url       = article.get("url", "#")
    tag_chips = "".join([
        f'<span class="tag-chip">{t}</span>' for t in tags
    ]) if tags else ""

    st.markdown(f"""
    <div class="news-card" style="animation-delay:{idx*0.08}s">
        {badge}
        <div class="card-title">{title}</div>
        <div class="card-meta">{source} &nbsp;&nbsp; {published}</div>
        <div class="score-bar-bg">
            <div class="score-bar-fill" style="width:{score_pct}%"></div>
        </div>
        <div style="font-size:0.75rem;margin-top:0.3rem;color:#888;">
            Urgency score: {score_pct}%
        </div>
        {f'<div class="card-reasoning">{reasoning}</div>' if reasoning else ''}
        {f'<div style="margin-top:0.8rem">{tag_chips}</div>' if tag_chips else ''}
    </div>
    """, unsafe_allow_html=True)

    with st.expander("Read more"):
        st.markdown(f"**Source:** {source}")
        st.markdown(f"**Published:** {published}")
        st.markdown(f"**Summary:** {article.get('summary', '')}")
        st.markdown(f"[Open full article]({url})")


# ─────────────────────────────────────────
# PAGINATION
# ─────────────────────────────────────────
def render_pagination(total: int, page_key: str):
    total_pages = max(1, -(-total // ARTICLES_PER_PAGE))
    if page_key not in st.session_state:
        st.session_state[page_key] = 1

    col_prev, col_info, col_next = st.columns([1, 3, 1])
    with col_prev:
        if st.button("Previous", key=f"{page_key}_prev",
                     disabled=st.session_state[page_key] <= 1):
            st.session_state[page_key] -= 1
            st.rerun()
    with col_info:
        st.markdown(
            f"<div style='text-align:center;padding:0.5rem;color:#6B6B6B;'>"
            f"Page {st.session_state[page_key]} of {total_pages}"
            f" ({total} total)</div>",
            unsafe_allow_html=True
        )
    with col_next:
        if st.button("Next", key=f"{page_key}_next",
                     disabled=st.session_state[page_key] >= total_pages):
            st.session_state[page_key] += 1
            st.rerun()

    start = (st.session_state[page_key] - 1) * ARTICLES_PER_PAGE
    return start, start + ARTICLES_PER_PAGE


# ─────────────────────────────────────────
# TAG SELECTOR
# ─────────────────────────────────────────
def render_tag_selector() -> list:
    st.markdown("#### Select up to 3 tags that matter to you today")
    all_tags = list(TAGS.keys())
    defaults = all_tags[:3]
    selected = []
    cols = st.columns(4)
    for i, tag in enumerate(all_tags):
        with cols[i % 4]:
            if st.checkbox(tag, value=tag in defaults, key=f"tag_{tag}"):
                selected.append(tag)
    if len(selected) > 3:
        st.warning("Max 3 tags. Please uncheck one to continue.")
        return selected[:3]
    return selected


# ─────────────────────────────────────────
# EXPORTS
# ─────────────────────────────────────────
def render_exports(articles: list):
    st.markdown("<div class='divider'></div>", unsafe_allow_html=True)
    st.markdown("#### Export Intelligence Report")

    col1, col2, col3 = st.columns(3)

    with col1:
        if st.button("Generate Excel", use_container_width=True):
            import openpyxl
            from config import PATHS
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Meridian Pulse"
            ws.append(["Title", "Source", "Published",
                        "Urgency", "Score", "Tags", "Reasoning", "URL"])
            for a in articles:
                ws.append([
                    a.get("title", ""), a.get("source", ""),
                    a.get("published", ""), a.get("urgency_label", ""),
                    a.get("urgency_score", 0),
                    ", ".join(a.get("tags_matched", [])),
                    a.get("reasoning", ""), a.get("url", ""),
                ])
            buf = io.BytesIO()
            wb.save(buf)
            buf.seek(0)
            st.download_button(
                "Download Excel",
                buf,
                file_name=f"meridian_{datetime.date.today()}.xlsx",
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                use_container_width=True
            )

    with col2:
        if st.button("Generate PDF", use_container_width=True):
            from fpdf import FPDF
            from config import PATHS
            pdf = FPDF()
            pdf.set_auto_page_break(auto=True, margin=15)
            pdf.add_page()
            pdf.set_font("Helvetica", "B", 20)
            pdf.cell(0, 12, "Meridian Pulse", ln=True, align="C")
            pdf.set_font("Helvetica", "", 10)
            pdf.cell(0, 8,
                f"Generated: {datetime.date.today()} | {len(articles)} articles",
                ln=True, align="C")
            pdf.ln(6)
            for a in articles[:50]:
                pdf.set_font("Helvetica", "B", 11)
                pdf.multi_cell(0, 7, a.get("title", "")[:90])
                pdf.set_font("Helvetica", "", 9)
                pdf.cell(0, 6,
                    f"{a.get('urgency_label','')} | "
                    f"Score: {int(a.get('urgency_score',0)*100)}% | "
                    f"{a.get('source','')} | {a.get('published','')}",
                    ln=True)
                if a.get("reasoning"):
                    pdf.multi_cell(0, 6, a["reasoning"][:200])
                pdf.ln(3)
            fname = f"meridian_{datetime.date.today()}.pdf"
            fpath = os.path.join(PATHS["pdf"], fname)
            pdf.output(fpath)
            with open(fpath, "rb") as f:
                st.download_button(
                    "Download PDF",
                    f, file_name=fname,
                    mime="application/pdf",
                    use_container_width=True
                )

    with col3:
        if st.button("Generate PPT", use_container_width=True):
            from pptx import Presentation
            from pptx.util import Inches, Pt
            from pptx.dml.color import RGBColor
            from config import PATHS
            prs = Presentation()
            prs.slide_width  = Inches(13.33)
            prs.slide_height = Inches(7.5)

            slide = prs.slides.add_slide(prs.slide_layouts[6])
            txb   = slide.shapes.add_textbox(Inches(1), Inches(2.5), Inches(11), Inches(2))
            tf    = txb.text_frame
            p     = tf.paragraphs[0]
            p.text = "Meridian Pulse"
            p.runs[0].font.size  = Pt(44)
            p.runs[0].font.bold  = True
            p.runs[0].font.color.rgb = RGBColor(0x1B, 0x4F, 0x8A)
            p.alignment = 1
            p2 = tf.add_paragraph()
            p2.text = f"Healthcare Intelligence Report  {datetime.date.today()}"
            p2.runs[0].font.size = Pt(18)
            p2.alignment = 1

            for a in [x for x in articles if x.get("urgency_label") == "🔴 High"][:20]:
                sl  = prs.slides.add_slide(prs.slide_layouts[6])
                tb  = sl.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(12), Inches(1.2))
                tff = tb.text_frame
                pp  = tff.paragraphs[0]
                pp.text = a.get("title", "")[:100]
                pp.runs[0].font.size  = Pt(22)
                pp.runs[0].font.bold  = True
                pp.runs[0].font.color.rgb = RGBColor(0x1B, 0x4F, 0x8A)
                tb2  = sl.shapes.add_textbox(Inches(0.5), Inches(1.8), Inches(12), Inches(3))
                tff2 = tb2.text_frame
                tff2.word_wrap = True
                pp2  = tff2.paragraphs[0]
                pp2.text = a.get("reasoning", "")[:400]
                pp2.runs[0].font.size = Pt(13)
                tb3  = sl.shapes.add_textbox(Inches(0.5), Inches(5.5), Inches(12), Inches(0.6))
                tff3 = tb3.text_frame
                pp3  = tff3.paragraphs[0]
                pp3.text = "Tags: " + ", ".join(a.get("tags_matched", []))
                pp3.runs[0].font.size  = Pt(11)
                pp3.runs[0].font.color.rgb = RGBColor(0x1B, 0x4F, 0x8A)

            buf = io.BytesIO()
            prs.save(buf)
            buf.seek(0)
            st.download_button(
                "Download PPT",
                buf,
                file_name=f"meridian_{datetime.date.today()}.pptx",
                mime="application/vnd.openxmlformats-officedocument.presentationml.presentation",
                use_container_width=True
            )


# ─────────────────────────────────────────
# SIDEBAR CHATBOT
# ─────────────────────────────────────────
def render_sidebar_chat():
    with st.sidebar:
        dark   = st.session_state.get("dark_mode", False)
        accent = "#4A90D9" if dark else "#1B4F8A"

        st.markdown(f"""
        <div style='padding:0.5rem 0 1rem;'>
            <div style='font-family:Playfair Display,serif;
                        font-size:1.2rem;font-weight:700;
                        color:{accent};margin-bottom:0.2rem;'>
                Ask Meridian
            </div>
            <div style='font-size:0.78rem;color:#9B9B9B;'>
                Your healthcare AI assistant
            </div>
        </div>
        """, unsafe_allow_html=True)

        articles = st.session_state.get("articles", [])
        if articles:
            st.caption(f"Connected to {len(articles)} scored articles")
            context = "\n".join([
                f"- [{a.get('urgency_label','')}] {a.get('title','')} "
                f"({a.get('source','')}, {a.get('published','')}): "
                f"{a.get('reasoning','')}"
                for a in articles[:20]
            ])
        else:
            st.caption("Fetch news on dashboard for full context")
            context = "No articles fetched yet."

        if "sidebar_chat" not in st.session_state:
            st.session_state["sidebar_chat"] = []

        if st.button("Clear chat", use_container_width=True, key="sidebar_clear"):
            st.session_state["sidebar_chat"] = []
            st.rerun()

        for msg in st.session_state["sidebar_chat"]:
            with st.chat_message(msg["role"]):
                st.markdown(msg["content"])

        query = st.chat_input("Ask anything...", key="sidebar_input")

        if query:
            st.session_state["sidebar_chat"].append({"role": "user", "content": query})
            history = "\n".join([
                f"{m['role'].upper()}: {m['content']}"
                for m in st.session_state["sidebar_chat"][-6:]
            ])
            with st.spinner("Thinking..."):
                try:
                    from engine.llm_handler import ask_llm, get_llm
                    from engine.prompts import CHATBOT_PROMPT
                    llm      = get_llm(role="scorer")
                    response = ask_llm(CHATBOT_PROMPT.format(
                        context=context,
                        history=history,
                        query=query,
                    ), llm)
                except Exception as e:
                    response = f"Meridian error: {e}"
            st.session_state["sidebar_chat"].append({"role": "assistant", "content": response})
            st.rerun()


# ─────────────────────────────────────────
# TAB 1 - URGENCY DASHBOARD
# ─────────────────────────────────────────
def tab_dashboard():
    selected_tags = render_tag_selector()

    if not selected_tags:
        st.info("Select at least 1 tag to load your dashboard.")
        return

    st.markdown("<div class='divider'></div>", unsafe_allow_html=True)

    if st.button("Fetch and Score News", type="primary", use_container_width=True):
        with st.spinner("Meridian Pulse is scanning the healthcare landscape..."):
            from ingestion.news_processor import fetch_all_news
            progress = st.progress(0)
            def update_progress(current, total):
                progress.progress(int((current / total) * 100))
            articles = fetch_all_news(
                selected_tags=selected_tags,
                progress_callback=update_progress,
            )
            st.session_state["articles"]      = articles
            st.session_state["selected_tags"] = selected_tags
            st.session_state["page"]          = 1
            progress.empty()

    articles = st.session_state.get("articles", [])

    if not articles:
        st.markdown("""
        <div style='text-align:center;padding:4rem;opacity:0.4;'>
            <div style='font-size:1.1rem;'>
                Select your tags and fetch news to begin
            </div>
        </div>
        """, unsafe_allow_html=True)
        return

    high   = sum(1 for a in articles if a["urgency_label"] == "🔴 High")
    medium = sum(1 for a in articles if a["urgency_label"] == "🟡 Medium")
    low    = sum(1 for a in articles if a["urgency_label"] == "🟢 Low")

    c1, c2, c3, c4 = st.columns(4)
    c1.metric("Total Articles", len(articles))
    c2.metric("High Priority",  high)
    c3.metric("Medium",         medium)
    c4.metric("Low",            low)

    st.markdown("<div class='divider'></div>", unsafe_allow_html=True)

    filter_label = st.selectbox(
        "Filter by urgency",
        ["All", "🔴 High", "🟡 Medium", "🟢 Low", "🔵 Minimal"]
    )
    filtered = articles if filter_label == "All" else [
        a for a in articles if a["urgency_label"] == filter_label
    ]
    filtered = filtered[:MAX_DASHBOARD_ARTICLES]

    start, end = render_pagination(len(filtered), "page")
    for i, article in enumerate(filtered[start:end]):
        render_card(article, i)

    if articles:
        render_exports(articles)


# ─────────────────────────────────────────
# TAB 2 - SEC INTELLIGENCE
# ─────────────────────────────────────────
def tab_sec():
    st.markdown("### SEC Filing Intelligence")
    st.markdown("Live competitor filings from UnitedHealth, Elevance, CVS, Humana, Cigna")
    st.markdown("<div class='divider'></div>", unsafe_allow_html=True)

    if st.button("Fetch SEC Filings", type="primary", use_container_width=True):
        with st.spinner("Pulling latest SEC filings..."):
            from ingestion.sec_fetcher import fetch_all_sec
            from engine.tag_scorer import score_all_articles
            selected_tags = st.session_state.get("selected_tags", list(TAGS.keys())[:3])
            filings = fetch_all_sec()
            scored  = score_all_articles(filings, selected_tags)
            st.session_state["sec_filings"] = scored
            st.session_state["sec_page"]    = 1

    filings = st.session_state.get("sec_filings", [])

    if not filings:
        st.markdown("""
        <div style='text-align:center;padding:4rem;opacity:0.4;'>
            <div style='font-size:1.1rem;'>
                Click fetch to load live SEC filings
            </div>
        </div>
        """, unsafe_allow_html=True)
        return

    col1, col2 = st.columns(2)
    with col1:
        company_filter = st.selectbox(
            "Filter by company",
            ["All"] + sorted(set(f["company"] for f in filings))
        )
    with col2:
        form_filter = st.selectbox(
            "Filter by filing type",
            ["All", "8-K", "10-K", "10-Q", "DEF 14A"]
        )

    filtered = filings
    if company_filter != "All":
        filtered = [f for f in filtered if f["company"] == company_filter]
    if form_filter != "All":
        filtered = [f for f in filtered if f["form_type"] == form_filter]

    start, end = render_pagination(len(filtered), "sec_page")

    for i, filing in enumerate(filtered[start:end]):
        render_card(filing, i)
        if st.button("Summarize with AI", key=f"sum_{filing['id']}", use_container_width=False):
            with st.spinner("Fetching and summarizing..."):
                try:
                    import requests
                    from bs4 import BeautifulSoup
                    from engine.llm_handler import ask_llm, get_llm
                    from config import SEC_HEADERS
                    doc_url = filing.get("doc_url", "")
                    r       = requests.get(doc_url, headers=SEC_HEADERS, timeout=15)
                    soup    = BeautifulSoup(r.text, "html.parser")
                    for tag in soup(["script", "style"]):
                        tag.decompose()
                    text    = soup.get_text(separator=" ", strip=True)[:6000]
                    llm     = get_llm(role="scorer")
                    prompt  = f"""You are a healthcare strategy analyst at BCBS.
Analyze this SEC filing and provide:
1. What happened (2-3 sentences)
2. Key financial figures mentioned
3. Strategic implications for BSC
4. Urgency level and why

Filing:
{text}"""
                    summary = ask_llm(prompt, llm)
                    st.markdown("#### AI Filing Summary")
                    st.markdown(summary)
                except Exception as e:
                    st.error(f"Could not summarize: {e}")


# ─────────────────────────────────────────
# TAB 3 - TRANSCRIPTS
# ─────────────────────────────────────────
def tab_transcripts():
    st.markdown("### Transcript Intelligence")
    st.markdown(
        "Select an earnings call transcript for an instant war room briefing."
    )
    st.markdown("<div class='divider'></div>", unsafe_allow_html=True)

    from config import PATHS
    transcript_dir = PATHS["transcripts"]
    existing_files = get_transcript_files(transcript_dir)

    content  = ""
    filename = ""

    col_select, col_upload = st.columns([2, 1])

    with col_select:
        if existing_files:
            selected = st.selectbox(
                "Select transcript",
                existing_files,
                help=f"From: {transcript_dir}"
            )
            filepath = os.path.join(transcript_dir, selected)
            content  = read_transcript(filepath)
            filename = selected
        else:
            st.info(f"No transcripts found in: {transcript_dir}")

    with col_upload:
        uploaded = st.file_uploader(
            "Upload new",
            type=["txt", "pdf", "docx"],
            label_visibility="visible"
        )
        if uploaded:
            save_path = os.path.join(transcript_dir, uploaded.name)
            with open(save_path, "wb") as f:
                f.write(uploaded.getbuffer())
            st.success(f"Saved: {uploaded.name}")
            st.rerun()

    if not content:
        st.markdown("""
        <div style='text-align:center;padding:3rem;opacity:0.4;'>
            <div style='font-size:1.1rem;'>
                Upload a transcript to begin intelligence briefing
            </div>
        </div>
        """, unsafe_allow_html=True)
        return

    st.caption(f"Transcript loaded: {len(content):,} characters")

    with st.expander("View raw transcript"):
        st.text(content[:3000])

    st.markdown("<div class='divider'></div>", unsafe_allow_html=True)

    if st.button("Run War Room Briefing", type="primary", use_container_width=True):
        with st.spinner("Meridian is analyzing the transcript..."):
            from engine.llm_handler import ask_llm, get_llm
            from engine.prompts import TRANSCRIPT_ANALYSIS_PROMPT
            llm = get_llm(role="scorer")
            try:
                raw    = ask_llm(TRANSCRIPT_ANALYSIS_PROMPT.format(
                    transcript=content[:5000]
                ), llm)
                raw    = raw.replace("```json", "").replace("```", "").strip()
                result = json.loads(raw)
                st.session_state["transcript_analysis"] = result
            except Exception as e:
                st.error(f"Analysis failed: {e}")

    analysis = st.session_state.get("transcript_analysis", {})

    if analysis:
        dark    = st.session_state.get("dark_mode", False)
        card_bg = "#1C1E26" if dark else "#FFFFFF"
        text_s  = "#9B9B9B" if dark else "#6B6B6B"
        border  = "#2E3140" if dark else "#E8E4DC"

        meta1, meta2, meta3 = st.columns(3)
        with meta1:
            st.markdown(f"**Company:** {analysis.get('company', 'Unknown')}")
        with meta2:
            st.markdown(f"**Period:** {analysis.get('period', 'Unknown')}")
        with meta3:
            speakers = analysis.get("speakers", [])
            st.markdown(f"**Speakers:** {len(speakers)}")

        st.markdown("<div class='divider'></div>", unsafe_allow_html=True)

        for card in [
            {"key": "pulse_check",    "title": "Pulse Check",    "color": "#1B4F8A", "type": "text"},
            {"key": "green_signals",  "title": "Green Signals",  "color": "#276749", "type": "list"},
            {"key": "red_signals",    "title": "Red Signals",    "color": "#CC0000", "type": "list"},
            {"key": "future_horizon", "title": "Future Horizon", "color": "#B8860B", "type": "list"},
            {"key": "data_spotlight", "title": "Data Spotlight", "color": "#6B3FA0", "type": "list"},
        ]:
            data = analysis.get(card["key"], [])
            if not data:
                continue
            if card["type"] == "text":
                content_html = f"<p style='color:{text_s};line-height:1.7;margin:0;'>{data}</p>"
            else:
                items = "".join([
                    f"<li style='color:{text_s};margin-bottom:0.4rem;'>{item}</li>"
                    for item in data
                ])
                content_html = f"<ul style='margin:0;padding-left:1.2rem;'>{items}</ul>"

            st.markdown(f"""
            <div style='background:{card_bg};border:1px solid {border};
                        border-radius:1rem;padding:1.5rem;margin-bottom:1rem;
                        border-left:4px solid {card["color"]};'>
                <div style='font-size:0.78rem;font-weight:600;
                            letter-spacing:1px;text-transform:uppercase;
                            color:{card["color"]};margin-bottom:0.8rem;'>
                    {card["title"]}
                </div>
                {content_html}
            </div>
            """, unsafe_allow_html=True)

        if speakers:
            with st.expander("View all speakers"):
                for s in speakers:
                    st.markdown(f"- {s}")

    st.markdown("<div class='divider'></div>", unsafe_allow_html=True)
    st.markdown("#### Ask anything about this transcript")

    ex_cols = st.columns(5)
    for i, ex in enumerate([
        "Who are the speakers?",
        "Summarize in 25 words",
        "What did the CFO say about MLR?",
        "List all financial figures",
        "What are the key risks?",
    ]):
        with ex_cols[i]:
            if st.button(ex, key=f"ex_{i}", use_container_width=True):
                st.session_state["transcript_query"]     = ex
                st.session_state["transcript_submitted"] = True

    query = st.text_input(
        "Your question",
        value=st.session_state.get("transcript_query", ""),
        placeholder="Ask anything about this transcript...",
    )

    if query != st.session_state.get("transcript_query", ""):
        st.session_state["transcript_query"]     = query
        st.session_state["transcript_submitted"] = False

    if st.button("Ask Meridian", type="primary"):
        st.session_state["transcript_submitted"] = True

    if st.session_state.get("transcript_submitted") and st.session_state.get("transcript_query"):
        current_query = st.session_state["transcript_query"]
        st.session_state["transcript_submitted"] = False
        with st.spinner("Analyzing..."):
            from engine.llm_handler import ask_llm, get_llm
            from engine.prompts import TRANSCRIPT_QA_PROMPT
            llm      = get_llm(role="scorer")
            response = ask_llm(TRANSCRIPT_QA_PROMPT.format(
                query=current_query,
                transcript=content[:5000]
            ), llm)
            st.session_state["transcript_response"] = response

    response = st.session_state.get("transcript_response", "")
    if response:
        for section in response.split("\n\n"):
            for label in ["QUERY", "RESPONSE", "CONFIDENCE", "SOURCE"]:
                if section.startswith(label):
                    st.markdown(
                        f"<div class='section-label'>{label.title()}</div>",
                        unsafe_allow_html=True
                    )
                    body = section.replace(label, "").strip()
                    if label == "SOURCE":
                        st.markdown(f"> {body}")
                    else:
                        st.markdown(body)
                    break

        dl1, dl2 = st.columns(2)
        with dl1:
            st.download_button(
                "Download TXT",
                response,
                file_name="meridian_transcript_insight.txt",
                mime="text/plain",
                use_container_width=True
            )
        with dl2:
            doc = __import__("docx").Document()
            doc.add_heading("Meridian Pulse", 0)
            doc.add_heading("Transcript Intelligence", 1)
            doc.add_paragraph(
                f"Query: {st.session_state.get('transcript_query', '')}"
            )
            doc.add_paragraph(response)
            buf = io.BytesIO()
            doc.save(buf)
            buf.seek(0)
            st.download_button(
                "Download Word Doc",
                buf,
                file_name=f"meridian_transcript_{datetime.date.today()}.docx",
                mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                use_container_width=True
            )


# ─────────────────────────────────────────
# TAB 3 EXTENDED - ASK MERIDIAN (full tab chatbot)
# ─────────────────────────────────────────
def tab_chatbot():
    st.markdown("### Ask Meridian")
    st.markdown(
        "Your AI intelligence assistant. Ask anything about healthcare news, "
        "competitors, BSC strategy, or generate reports."
    )
    st.markdown("<div class='divider'></div>", unsafe_allow_html=True)

    dark    = st.session_state.get("dark_mode", False)
    card_bg = "#1C1E26" if dark else "#FFFFFF"
    border  = "#2E3140" if dark else "#E8E4DC"
    text_s  = "#9B9B9B" if dark else "#6B6B6B"
    accent  = "#4A90D9" if dark else "#1B4F8A"

    articles = st.session_state.get("articles", [])
    if articles:
        st.caption(f"Connected to {len(articles)} scored articles from dashboard")
        context = "\n".join([
            f"- [{a.get('urgency_label','')}] {a.get('title','')} "
            f"({a.get('source','')}, {a.get('published','')}): "
            f"{a.get('reasoning','')}"
            for a in articles[:30]
        ])
    else:
        st.info("Tip: Fetch news on the Urgency Dashboard tab first to give Meridian full context.")
        context = "No articles fetched yet."

    st.markdown("<div class='divider'></div>", unsafe_allow_html=True)

    # Suggestions
    st.markdown(
        f"<div style='font-size:0.8rem;color:{text_s};margin-bottom:0.5rem;'>"
        f"Try asking:</div>",
        unsafe_allow_html=True
    )
    sug_cols = st.columns(3)
    for i, sug in enumerate([
        "What are the biggest regulatory risks this week?",
        "Summarize competitor moves in Medicare Advantage",
        "What should BSC watch out for?",
        "Generate a report of top 5 high priority articles",
        "What women in healthcare news is trending?",
        "Export this week's news as Excel",
    ]):
        with sug_cols[i % 3]:
            if st.button(sug, key=f"sug_{i}", use_container_width=True):
                st.session_state["chat_input"]     = sug
                st.session_state["chat_submitted"] = True

    st.markdown("<div class='divider'></div>", unsafe_allow_html=True)

    if "chat_history" not in st.session_state:
        st.session_state["chat_history"] = []

    # Chat history
    for msg in st.session_state["chat_history"]:
        if msg["role"] == "user":
            st.markdown(
                f'<div class="chat-bubble-user">{msg["content"]}</div>',
                unsafe_allow_html=True
            )
        else:
            st.markdown(
                f'<div class="chat-bubble-bot">{msg["content"]}</div>',
                unsafe_allow_html=True
            )
            # Export row after each response
            e1, e2, e3, e4 = st.columns(4)
            with e1:
                st.download_button(
                    "TXT",
                    msg["content"],
                    file_name="meridian_response.txt",
                    mime="text/plain",
                    use_container_width=True,
                    key=f"txt_{msg.get('id','')}"
                )
            with e2:
                if st.button("Excel", use_container_width=True, key=f"excel_{msg.get('id','')}"):
                    import openpyxl
                    from config import PATHS
                    wb = openpyxl.Workbook()
                    ws = wb.active
                    ws.append(["Query", "Response", "Date"])
                    ws.append([msg.get("query",""), msg["content"], str(datetime.date.today())])
                    if articles:
                        ws2 = wb.create_sheet("Articles")
                        ws2.append(["Title","Source","Published","Urgency","Score","Tags"])
                        for a in articles:
                            ws2.append([a.get("title",""),a.get("source",""),
                                       a.get("published",""),a.get("urgency_label",""),
                                       a.get("urgency_score",0),
                                       ", ".join(a.get("tags_matched",[]))])
                    buf = io.BytesIO()
                    wb.save(buf)
                    buf.seek(0)
                    st.download_button(
                        "Download Excel",
                        buf,
                        file_name=f"meridian_{datetime.date.today()}.xlsx",
                        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                        key=f"dl_xl_{msg.get('id','')}"
                    )
            with e3:
                if st.button("Word", use_container_width=True, key=f"word_{msg.get('id','')}"):
                    from docx import Document
                    from config import PATHS
                    doc = Document()
                    doc.add_heading("Meridian Pulse", 0)
                    doc.add_paragraph(f"Query: {msg.get('query','')}")
                    doc.add_paragraph(msg["content"])
                    buf = io.BytesIO()
                    doc.save(buf)
                    buf.seek(0)
                    st.download_button(
                        "Download Word",
                        buf,
                        file_name=f"meridian_{datetime.date.today()}.docx",
                        mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                        key=f"dl_wd_{msg.get('id','')}"
                    )
            with e4:
                if st.button("PPT", use_container_width=True, key=f"ppt_{msg.get('id','')}"):
                    from pptx import Presentation
                    from pptx.util import Inches, Pt
                    from pptx.dml.color import RGBColor
                    prs  = Presentation()
                    sl   = prs.slides.add_slide(prs.slide_layouts[6])
                    txb  = sl.shapes.add_textbox(Inches(0.5),Inches(0.5),Inches(12),Inches(6))
                    tf   = txb.text_frame
                    tf.word_wrap = True
                    p    = tf.paragraphs[0]
                    p.text = msg["content"][:800]
                    p.runs[0].font.size = Pt(14)
                    buf  = io.BytesIO()
                    prs.save(buf)
                    buf.seek(0)
                    st.download_button(
                        "Download PPT",
                        buf,
                        file_name=f"meridian_{datetime.date.today()}.pptx",
                        mime="application/vnd.openxmlformats-officedocument.presentationml.presentation",
                        key=f"dl_ppt_{msg.get('id','')}"
                    )

    # Input
    user_input = st.text_input(
        "Message Meridian",
        value=st.session_state.get("chat_input", ""),
        placeholder="Ask about healthcare news, competitors, BSC strategy...",
        key="chat_text_input"
    )

    c1, c2 = st.columns([3, 1])
    with c1:
        send = st.button("Send", type="primary", use_container_width=True)
    with c2:
        if st.button("Clear chat", use_container_width=True):
            st.session_state["chat_history"] = []
            st.session_state["chat_input"]   = ""
            st.rerun()

    query     = user_input or st.session_state.get("chat_input", "")
    submitted = send or st.session_state.get("chat_submitted", False)

    if submitted and query:
        st.session_state["chat_submitted"] = False
        st.session_state["chat_input"]     = ""

        st.session_state["chat_history"].append({
            "role": "user", "content": query,
            "id":   len(st.session_state["chat_history"]),
        })

        history = "\n".join([
            f"{m['role'].upper()}: {m['content']}"
            for m in st.session_state["chat_history"][-6:]
        ])

        with st.spinner("Meridian is thinking..."):
            from engine.llm_handler import ask_llm, get_llm
            from engine.prompts import CHATBOT_PROMPT
            llm      = get_llm(role="scorer")
            response = ask_llm(CHATBOT_PROMPT.format(
                context=context,
                history=history,
                query=query,
            ), llm)

        st.session_state["chat_history"].append({
            "role": "assistant", "content": response,
            "query": query,
            "id":    len(st.session_state["chat_history"]),
        })
        st.rerun()


# ─────────────────────────────────────────
# TAB 4 - BEHIND MERIDIAN PULSE
# ─────────────────────────────────────────

def tab_hire():
    import base64

    qr_path = os.path.join(os.path.dirname(__file__), "assets", "linkedin.png")

    try:
        with open(qr_path, "rb") as f:
            qr_b64 = base64.b64encode(f.read()).decode()
            qr_html = f"<img src='data:image/png;base64,{qr_b64}' style='border-radius:1rem;margin-top:1rem;width:200px;'/>"
    except FileNotFoundError:
        qr_html = "<div style='color:#9B9B9B;font-size:0.8rem;'>QR image not found</div>"


    st.markdown(f"""
    <div class='hire-card'>
        <div class='hire-name'>Akansha Rawat</div>
        <div class='hire-role'>Healthcare AI and Strategy</div>
        <div class='divider'></div>
        <p style='color:#6B6B6B;font-size:0.95rem;line-height:1.8;text-align:center;max-width:480px;margin:0 auto 1.5rem;'>
            Meridian Pulse was built to solve a real problem.
            BSC analysts spending 4 hours daily on manual news research.
            This tool brings that down to minutes using AI powered
            urgency scoring, live data pipelines, and intelligent filtering.
        </p>
        <div class='divider'></div>
        <a class='contact-link' href='mailto:amrawat@uci.edu'>
            Email Here
        </a>
        <a class='contact-link'
           href='https://www.linkedin.com/in/akansharawat01/'
           target='_blank'>
            LinkedIn Here 
        </a>
        <div class='divider'></div>
        {qr_html}
        <div style='font-size:0.8rem;color:#9B9B9B;margin-top:0.8rem;'>
            Scan to connect on LinkedIn
        </div>
    </div>
    """, unsafe_allow_html=True)


# ─────────────────────────────────────────
# MAIN
# ─────────────────────────────────────────
def main():
    if "dark_mode" not in st.session_state:
        st.session_state["dark_mode"] = False

    col1, col2 = st.columns([10, 1])
    with col2:
        dark = st.toggle("Night", value=st.session_state["dark_mode"])
        st.session_state["dark_mode"] = dark

    load_styles(dark)
    render_sidebar_chat()
    render_hero()

    tab1, tab2, tab3, tab4 = st.tabs([
        "Urgency Dashboard",
        "SEC Intelligence",
        "Transcripts",
        "Behind Meridian Pulse",
    ])

    with tab1:
        tab_dashboard()
    with tab2:
        tab_sec()
    with tab3:
        tab_transcripts()
    with tab4:
        tab_hire()


if __name__ == "__main__":
    main()
