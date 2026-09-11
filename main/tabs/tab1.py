# TAB 1 - URGENCY DASHBOARD

import os
import sys
import io
import datetime
import streamlit as st

ROOT = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
sys.path.append(ROOT)
sys.path.append(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from config import TAGS, MAX_DASHBOARD_ARTICLES, PATHS
from components import render_card, render_pagination


def render_tag_selector() -> list:
    st.markdown("#### Select up to 3 tags that matter to you today")
    all_tags = list(TAGS.keys())
    defaults = all_tags[:3]
    selected = []
    cols = st.columns(4)
    for i, tag in enumerate(all_tags):
        with cols[i % 4]:
            if st.checkbox(tag, value=tag in defaults, key=f"tag_{tag}"):
                selected.append(tag)
    if len(selected) > 3:
        st.warning("Max 3 tags. Please uncheck one to continue.")
        return selected[:3]
    return selected


def render_exports(articles: list):
    st.markdown("<div class='divider'></div>", unsafe_allow_html=True)
    st.markdown("#### Export Intelligence Report")

    col1, col2, col3 = st.columns(3)

    with col1:
        if st.button("Generate Excel", use_container_width=True):
            import openpyxl
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Meridian Pulse"
            ws.append(["Title", "Source", "Published", "Urgency", "Score", "Tags", "Reasoning", "URL"])
            for a in articles:
                ws.append([
                    a.get("title", ""), a.get("source", ""),
                    a.get("published", ""), a.get("urgency_label", ""),
                    a.get("urgency_score", 0),
                    ", ".join(a.get("tags_matched", [])),
                    a.get("reasoning", ""), a.get("url", ""),
                ])
            buf = io.BytesIO()
            wb.save(buf)
            buf.seek(0)
            st.download_button(
                "Download Excel", buf,
                file_name=f"meridian_{datetime.date.today()}.xlsx",
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                use_container_width=True
            )

    with col2:
        if st.button("Generate PDF", use_container_width=True):
            from fpdf import FPDF
            pdf = FPDF()
            pdf.set_auto_page_break(auto=True, margin=15)
            pdf.add_page()
            pdf.set_font("Helvetica", "B", 20)
            pdf.cell(0, 12, "Meridian Pulse", ln=True, align="C")
            pdf.set_font("Helvetica", "", 10)
            pdf.cell(0, 8, f"Generated: {datetime.date.today()} | {len(articles)} articles", ln=True, align="C")
            pdf.ln(6)
            for a in articles[:50]:
                pdf.set_font("Helvetica", "B", 11)
                pdf.multi_cell(0, 7, a.get("title", "")[:90])
                pdf.set_font("Helvetica", "", 9)
                pdf.cell(0, 6,
                    f"{a.get('urgency_label','')} | Score: {int(a.get('urgency_score',0)*100)}% | "
                    f"{a.get('source','')} | {a.get('published','')}",
                    ln=True)
                if a.get("reasoning"):
                    pdf.multi_cell(0, 6, a["reasoning"][:200])
                pdf.ln(3)
            fname = f"meridian_{datetime.date.today()}.pdf"
            fpath = os.path.join(PATHS["pdf"], fname)
            pdf.output(fpath)
            with open(fpath, "rb") as f:
                st.download_button(
                    "Download PDF", f, file_name=fname,
                    mime="application/pdf", use_container_width=True
                )

    with col3:
        if st.button("Generate PPT", use_container_width=True):
            from pptx import Presentation
            from pptx.util import Inches, Pt
            from pptx.dml.color import RGBColor
            prs = Presentation()
            prs.slide_width  = Inches(13.33)
            prs.slide_height = Inches(7.5)
            slide = prs.slides.add_slide(prs.slide_layouts[6])
            txb   = slide.shapes.add_textbox(Inches(1), Inches(2.5), Inches(11), Inches(2))
            tf    = txb.text_frame
            p     = tf.paragraphs[0]
            p.text = "Meridian Pulse"
            p.runs[0].font.size  = Pt(44)
            p.runs[0].font.bold  = True
            p.runs[0].font.color.rgb = RGBColor(0x1B, 0x4F, 0x8A)
            p.alignment = 1
            p2 = tf.add_paragraph()
            p2.text = f"Healthcare Intelligence Report  {datetime.date.today()}"
            p2.runs[0].font.size = Pt(18)
            p2.alignment = 1
            for a in [x for x in articles if x.get("urgency_label") == "🔴 High"][:20]:
                sl  = prs.slides.add_slide(prs.slide_layouts[6])
                tb  = sl.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(12), Inches(1.2))
                pp  = tb.text_frame.paragraphs[0]
                pp.text = a.get("title", "")[:100]
                pp.runs[0].font.size  = Pt(22)
                pp.runs[0].font.bold  = True
                pp.runs[0].font.color.rgb = RGBColor(0x1B, 0x4F, 0x8A)
                tb2 = sl.shapes.add_textbox(Inches(0.5), Inches(1.8), Inches(12), Inches(3))
                tb2.text_frame.word_wrap = True
                pp2 = tb2.text_frame.paragraphs[0]
                pp2.text = a.get("reasoning", "")[:400]
                pp2.runs[0].font.size = Pt(13)
                tb3 = sl.shapes.add_textbox(Inches(0.5), Inches(5.5), Inches(12), Inches(0.6))
                pp3 = tb3.text_frame.paragraphs[0]
                pp3.text = "Tags: " + ", ".join(a.get("tags_matched", []))
                pp3.runs[0].font.size  = Pt(11)
                pp3.runs[0].font.color.rgb = RGBColor(0x1B, 0x4F, 0x8A)
            buf = io.BytesIO()
            prs.save(buf)
            buf.seek(0)
            st.download_button(
                "Download PPT", buf,
                file_name=f"meridian_{datetime.date.today()}.pptx",
                mime="application/vnd.openxmlformats-officedocument.presentationml.presentation",
                use_container_width=True
            )


def tab_dashboard():
    selected_tags = render_tag_selector()

    if not selected_tags:
        st.info("Select at least 1 tag to load your dashboard.")
        return

    st.markdown("<div class='divider'></div>", unsafe_allow_html=True)

    if st.button("Fetch and Score News", type="primary", use_container_width=True):
        with st.spinner("Meridian Pulse is scanning the healthcare landscape..."):
            from ingestion.news_processor import fetch_all_news
            progress = st.progress(0)
            def update_progress(current, total):
                progress.progress(int((current / total) * 100))
            articles = fetch_all_news(
                selected_tags=selected_tags,
                progress_callback=update_progress,
            )
            st.session_state["articles"]      = articles
            st.session_state["selected_tags"] = selected_tags
            st.session_state["page"]          = 1
            progress.empty()

    articles = st.session_state.get("articles", [])

    if not articles:
        st.markdown("""
        <div style='text-align:center;padding:4rem;opacity:0.4;'>
            <div style='font-size:1.1rem;'>Select your tags and fetch news to begin</div>
        </div>
        """, unsafe_allow_html=True)
        return

    high   = sum(1 for a in articles if a["urgency_label"] == "🔴 High")
    medium = sum(1 for a in articles if a["urgency_label"] == "🟡 Medium")
    low    = sum(1 for a in articles if a["urgency_label"] == "🟢 Low")

    c1, c2, c3, c4 = st.columns(4)
    c1.metric("Total Articles", len(articles))
    c2.metric("High Priority",  high)
    c3.metric("Medium",         medium)
    c4.metric("Low",            low)

    st.markdown("<div class='divider'></div>", unsafe_allow_html=True)

    filter_label = st.selectbox("Filter by urgency",
        ["All", "🔴 High", "🟡 Medium", "🟢 Low", "🔵 Minimal"])
    filtered = articles if filter_label == "All" else [
        a for a in articles if a["urgency_label"] == filter_label
    ]
    filtered = filtered[:MAX_DASHBOARD_ARTICLES]

    start, end = render_pagination(len(filtered), "page")
    for i, article in enumerate(filtered[start:end]):
        render_card(article, i)

    if articles:
        render_exports(articles)
